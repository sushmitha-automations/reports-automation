"""
High-Performance Data Loaders using Polars exclusively (Zero Pandas).
Handles 3-way user identity resolution, facility mapping, biometric punch loading,
and onboarding funnel data hygiene.
"""
import os
import re
from datetime import datetime, date
from typing import Dict, List, Tuple, Any, Optional
import polars as pl
import openpyxl

from config import LM_PODS, HL_REGIONS, FM_REGIONS

# ==============================================================================
# STRING & POD NORMALIZATION HELPERS
# ==============================================================================

def normalize_pod(raw_pod: Optional[str], branch_code: Optional[str]) -> str:
    p = str(raw_pod or '').strip()
    p_lower = p.lower()
    if p_lower == 'pod_bom/pnq':
        return 'POD_BOM/PNQ'
    if p in ('POD_ROK', 'POD_KL', 'POD_ROK/KL'):
        return 'POD_ROK/KL'
    if p in ('POD_NCR', 'POD_NCR/UP_West', 'POD_UP_West'):
        return 'POD_NCR/UP_West'
    if not p and 'chn_central_large' in str(branch_code or '').lower():
        return 'POD_CHN'
    return p

def map_hl_pod(raw_pod: Optional[str], state: Optional[str]) -> str:
    p = str(raw_pod or '').upper()
    s = str(state or '').upper()
    if any(k in p for k in ['NCR', 'NORTH', 'NROI']) or any(k in s for k in ['DELHI', 'PUNJAB', 'HARYANA', 'UTTAR PRADESH']):
        return 'HL_North'
    if any(k in p for k in ['WROI', 'WEST', 'BOM', 'PNQ']) or any(k in s for k in ['MAHARASHTRA', 'GUJARAT']):
        return 'HL_West'
    if any(k in p for k in ['EROI', 'EAST', 'KOLKATA']) or any(k in s for k in ['WEST BENGAL', 'BIHAR', 'ODISHA', 'ASSAM']):
        return 'HL_East'
    if any(k in p for k in ['SROI', 'SOUTH', 'HYD', 'BLR', 'CHENNAI']) or any(k in s for k in ['KARNATAKA', 'TAMIL NADU', 'TELANGANA', 'ANDHRA']):
        return 'HL_South'
    return 'HL_North'

def map_fm_pod(branch_code: Optional[str], state: Optional[str]) -> str:
    b = str(branch_code or '').upper()
    s = str(state or '').upper()
    if any(k in s for k in ['WEST BENGAL', 'BIHAR', 'ODISHA', 'JHARKHAND', 'ASSAM']) or any(k in b for k in ['CCU', 'PAT', 'NE']):
        return 'FM_East'
    if any(k in s for k in ['TAMIL NADU', 'KARNATAKA', 'TELANGANA', 'ANDHRA', 'KERALA']) or any(k in b for k in ['BLR', 'HYD', 'CHN']):
        return 'FM_South'
    if any(k in s for k in ['MAHARASHTRA', 'GUJARAT', 'RAJASTHAN', 'MADHYA PRADESH']) or any(k in b for k in ['BOM', 'PNQ', 'ST']):
        return 'FM_West'
    return 'FM_North'

def clean_digits(val: Any) -> str:
    """Strips non-digits safely."""
    if val is None:
        return ''
    return re.sub(r'\D', '', str(val))

def extract_numeric_id(val: Any) -> Optional[int]:
    """Extracts numeric id from values like 'SFXPROD1234' or 1234."""
    if val is None:
        return None
    s = str(val).strip()
    match = re.search(r'\d+', s)
    return int(match.group(0)) if match else None

# ==============================================================================
# POLARS BASE LOADERS
# ==============================================================================

def load_facility_master(csv_path: str) -> Tuple[List[Dict[str, Any]], Dict[str, str]]:
    """Loads Valinor Dump - Facility.csv using Polars."""
    if not os.path.exists(csv_path):
        return [], {}

    # Read schema all as string to avoid type inference issues
    df = pl.read_csv(csv_path, infer_schema_length=0, truncate_ragged_lines=True)
    
    # Standardize column names
    col_map = {c: c.strip().lower() for c in df.columns}
    df = df.rename(col_map)

    id_col = next((c for c in df.columns if 'facility id' in c or c == 'id'), df.columns[0])
    name_col = next((c for c in df.columns if 'facility name' in c or 'name' in c), df.columns[1])
    type_col = next((c for c in df.columns if 'type' in c), 'type')
    saruman_col = next((c for c in df.columns if 'saruman' in c), 'saruman id')
    darsa_col = next((c for c in df.columns if 'darsa' in c), 'darsa id')
    pod_name_col = next((c for c in df.columns if 'pod name' in c), 'pod name')
    pod_id_col = next((c for c in df.columns if 'pod id' in c), 'pod id')
    state_name_col = next((c for c in df.columns if 'state name' in c), 'state name')
    state_id_col = next((c for c in df.columns if 'state id' in c), 'state id')

    facilities = []
    facility_pod_map = {}

    for row in df.iter_rows(named=True):
        f_name = str(row.get(name_col) or '').strip()
        p_name = str(row.get(pod_name_col) or '').strip()
        fac_obj = {
            'id': row.get(id_col) or '',
            'name': f_name,
            'type': row.get(type_col) or '',
            'saruman_id': row.get(saruman_col) or '',
            'darsa_id': row.get(darsa_col) or '',
            'pod_name': p_name,
            'pod_id': row.get(pod_id_col) or '',
            'state_name': row.get(state_name_col) or '',
            'state_id': row.get(state_id_col) or ''
        }
        facilities.append(fac_obj)
        if f_name:
            facility_pod_map[f_name.lower()] = p_name

    return facilities, facility_pod_map

def load_user_master(
    csv_paths: List[str],
    least_priority_excel_path: Optional[str],
    facility_pod_map: Dict[str, str]
) -> Tuple[List[Dict[str, Any]], Dict[str, Any], Dict[str, Any], Dict[int, Any]]:
    """
    Loads all_user_data CSVs using Polars, builds fast identity lookups
    (Aadhaar 12-digit, Mobile 10-digit, Numeric ID).
    """
    val_map_by_aadh: Dict[str, Any] = {}
    val_map_by_mob: Dict[str, Any] = {}
    val_map_by_id: Dict[int, Any] = {}
    user_rows: List[Dict[str, Any]] = []

    for path in csv_paths:
        if not os.path.exists(path):
            continue
        
        # Read with Polars ensuring string types for sensitive fields
        df = pl.read_csv(
            path,
            infer_schema_length=0,
            truncate_ragged_lines=True,
            ignore_errors=True
        )
        
        # Map columns flexibly
        cols = {c.strip().lower().replace(' ', '_'): c for c in df.columns}
        
        id_col = cols.get('id')
        name_col = cols.get('name')
        status_col = cols.get('status')
        phone_col = cols.get('phone_number')
        aadh_col = cols.get('aadhaar_number')
        fac_col = cols.get('registered_facility') or cols.get('facility_name')
        created_col = cols.get('created_date_ist') or cols.get('registered_date')
        role_type_col = cols.get('role_type')
        rolename_col = cols.get('rolename') or cols.get('role')

        for row in df.iter_rows(named=True):
            raw_id = row.get(id_col)
            num_id = extract_numeric_id(raw_id)
            if not num_id or num_id in val_map_by_id:
                continue

            phone_raw = str(row.get(phone_col) or '').strip()
            aadh_raw = str(row.get(aadh_col) or '').strip()
            fac_name = str(row.get(fac_col) or '').strip()
            role_name = str(row.get(rolename_col) or '').strip()

            mob_clean = clean_digits(phone_raw)[-10:]
            aadh_clean = clean_digits(aadh_raw)[-12:]
            pod = facility_pod_map.get(fac_name.lower(), '')

            rec = {
                'numericId': num_id,
                'phone_number': phone_raw,
                'name': str(row.get(name_col) or '').strip(),
                'status': str(row.get(status_col) or '').strip(),
                'aadhaar_number': aadh_raw,
                'registered_facility': fac_name,
                'created_date_ist': str(row.get(created_col) or '').strip(),
                'role_type': str(row.get(role_type_col) or '').strip(),
                'rolename': role_name,
                'facility_name': fac_name,
                'pod': pod
            }

            user_rows.append(rec)
            val_map_by_id[num_id] = rec
            if len(mob_clean) == 10 and mob_clean not in val_map_by_mob:
                val_map_by_mob[mob_clean] = rec
            if len(aadh_clean) == 12 and aadh_clean not in val_map_by_aadh:
                val_map_by_aadh[aadh_clean] = rec

    # Fallback to Least Priority Base - Valinor.xlsx if provided
    if least_priority_excel_path and os.path.exists(least_priority_excel_path):
        wb = openpyxl.load_workbook(least_priority_excel_path, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(c.value or '').strip().lower() for c in next(ws.iter_rows(min_row=1, max_row=1))]
        
        id_idx = next((i for i, h in enumerate(headers) if h in ('id', 'val id')), 0)
        phone_idx = next((i for i, h in enumerate(headers) if 'phone' in h), 2)
        name_idx = next((i for i, h in enumerate(headers) if 'name' in h), 3)
        status_idx = next((i for i, h in enumerate(headers) if 'status' in h), 4)
        aadh_idx = next((i for i, h in enumerate(headers) if 'aadhaar' in h), 5)
        fac_idx = next((i for i, h in enumerate(headers) if 'facility' in h), 6)
        created_idx = next((i for i, h in enumerate(headers) if 'created' in h), 7)

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[id_idx]:
                continue
            num_id = extract_numeric_id(row[id_idx])
            if not num_id or num_id in val_map_by_id:
                continue

            phone_raw = str(row[phone_idx] or '').strip()
            aadh_raw = str(row[aadh_idx] or '').strip()
            fac_name = str(row[fac_idx] or '').strip()
            mob_clean = clean_digits(phone_raw)[-10:]
            aadh_clean = clean_digits(aadh_raw)[-12:]
            pod = facility_pod_map.get(fac_name.lower(), '')

            rec = {
                'numericId': num_id,
                'phone_number': phone_raw,
                'name': str(row[name_idx] or '').strip(),
                'status': str(row[status_idx] or '').strip(),
                'aadhaar_number': aadh_raw,
                'registered_facility': fac_name,
                'created_date_ist': str(row[created_idx] or '').strip(),
                'role_type': 'DE',
                'rolename': 'DE',
                'facility_name': fac_name,
                'pod': pod
            }
            user_rows.append(rec)
            val_map_by_id[num_id] = rec
            if len(mob_clean) == 10 and mob_clean not in val_map_by_mob:
                val_map_by_mob[mob_clean] = rec
            if len(aadh_clean) == 12 and aadh_clean not in val_map_by_aadh:
                val_map_by_aadh[aadh_clean] = rec
        wb.close()

    return user_rows, val_map_by_aadh, val_map_by_mob, val_map_by_id

def load_valinor_attendance(file_path: str, date_list: List[str]) -> Tuple[List[Dict[str, Any]], Dict[int, Dict[str, str]]]:
    """
    Loads Valinor Attendance.csv using Polars and indexes daily attendance punches by numeric ID.
    Normalizes 'P', '1', 'P1' -> 'P', 'Tech Issue' -> 'Tech Issue', else 'A'.
    """
    if not os.path.exists(file_path):
        return [], {}

    att_map: Dict[int, Dict[str, str]] = {}
    att_rows: List[Dict[str, Any]] = []

    if file_path.endswith('.csv'):
        df = pl.read_csv(file_path, infer_schema_length=0, truncate_ragged_lines=True)
        cols = df.columns
        emp_col = cols[1]
        name_col = cols[0]
        fac_col = cols[2]

        date_col_map = {}
        for d in date_list:
            # Handle DD-MM-YYYY vs YYYY-MM-DD
            d_parts = d.split('-')
            d_alt = f"{d_parts[2]}-{d_parts[1]}-{d_parts[0]}" if len(d_parts) == 3 else d
            matched_col = next((c for c in cols if c == d or c == d_alt or c.endswith(d[-5:])), None)
            if matched_col:
                date_col_map[d] = matched_col

        for row in df.iter_rows(named=True):
            raw_emp = row.get(emp_col)
            num_id = extract_numeric_id(raw_emp)
            if not num_id:
                continue

            days = {}
            for d in date_list:
                col_name = date_col_map.get(d)
                raw_val = str(row.get(col_name) or '').strip().upper() if col_name else 'A'
                if raw_val in ('P', '1', 'P1'):
                    val = 'P'
                elif 'TECH' in raw_val:
                    val = 'Tech Issue'
                else:
                    val = 'A'
                days[d] = val

            att_obj = {
                'empId': num_id,
                'name': str(row.get(name_col) or '').strip(),
                'facility': str(row.get(fac_col) or '').strip(),
                'days': days
            }
            att_rows.append(att_obj)
            att_map[num_id] = days

    return att_rows, att_map

def load_chroma_masters(
    file_configs: List[Dict[str, str]],
    date_list: List[str],
    facility_pod_map: Dict[str, str]
) -> List[Dict[str, Any]]:
    """
    Loads Chroma Master workbooks (LM, FM, HL & Corp) using openpyxl,
    extracts metadata and 30-day attendance punches.
    """
    chroma_rows: List[Dict[str, Any]] = []
    seen_emp_ids = set()

    for item in file_configs:
        fp = item['file_path']
        f_type = item['type']
        if not os.path.exists(fp):
            continue

        wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
        ws = wb.active

        for r_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[0]:
                continue
            emp_id = str(row[0]).strip()
            if emp_id in seen_emp_ids:
                continue
            seen_emp_ids.add(emp_id)

            val_id_raw = extract_numeric_id(row[1]) if len(row) > 1 and row[1] else None
            mob_str = clean_digits(row[7])[-10:] if len(row) > 7 else ''
            aadh_str = clean_digits(row[8])[-12:] if len(row) > 8 else ''
            branch_code = str(row[14] or '').strip() if len(row) > 14 else ''
            state = str(row[17] or '').strip() if len(row) > 17 else ''
            base_pod_name = str(row[23] or '').strip() if len(row) > 23 else ''

            # Day punches (Day 1 to 30)
            date_punches = []
            for d in range(1, len(date_list) + 1):
                col_idx = 23 + d
                raw_punch = row[col_idx] if len(row) > col_idx and row[col_idx] is not None else 'NA'
                s_punch = str(raw_punch).strip().upper()
                if s_punch.startswith('#'):
                    s_punch = 'NA'
                elif s_punch in ('P1',):
                    s_punch = 'P'
                elif s_punch in ('0', 'LEFT', 'NULL'):
                    s_punch = 'A'
                elif not s_punch:
                    s_punch = 'NA'
                date_punches.append(s_punch)

            chroma_rows.append({
                'emp_id': emp_id,
                'valinor_id_raw': val_id_raw,
                'name': str(row[2] or '').strip() if len(row) > 2 else '',
                'doj': str(row[3] or '') if len(row) > 3 else '',
                'dob': str(row[4] or '') if len(row) > 4 else '',
                'age': str(row[5] or '') if len(row) > 5 else '',
                'father_name': str(row[6] or '').strip() if len(row) > 6 else '',
                'mobile': mob_str,
                'aadhaar': aadh_str,
                'pan': str(row[9] or '').strip() if len(row) > 9 else '',
                'gender': str(row[10] or '').strip() if len(row) > 10 else '',
                'skill': str(row[11] or '').strip() if len(row) > 11 else '',
                'designation': str(row[12] or '').strip() if len(row) > 12 else '',
                'payroll_source': str(row[13] or '').strip() if len(row) > 13 else '',
                'branch_code': branch_code,
                'cost_center': str(row[15] or '').strip() if len(row) > 15 else '',
                'sub_cost_center': str(row[16] or '').strip() if len(row) > 16 else '',
                'state': state,
                'city': str(row[18] or '').strip() if len(row) > 18 else '',
                'vendor_name': str(row[19] or '').strip() if len(row) > 19 else '',
                'rm_emp_id': str(row[20] or '').strip() if len(row) > 20 else '',
                'rm_name': str(row[21] or '').strip() if len(row) > 21 else '',
                'rm_status': str(row[22] or '').strip() if len(row) > 22 else '',
                'base_pod_name': base_pod_name,
                'file_type': f_type,
                'date_punches': date_punches,
                'raw_row': list(row)
            })
        wb.close()

    return chroma_rows

def load_onboarding_funnel(
    funnel_csv_path: str,
    migration_excel_path: Optional[str],
    start_date_str: str,
    cutoff_date_str: str
) -> Tuple[List[str], List[List[Any]], List[str]]:
    """
    Loads onboarding funnel using Polars, excludes migrated employees,
    filters by 'Onboarded*' status and date range, calculates Ageing.
    """
    migrated_ids = set()
    if migration_excel_path and os.path.exists(migration_excel_path):
        wb = openpyxl.load_workbook(migration_excel_path, read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) > 6 and row[6]:
                val_id = str(row[6]).strip().upper()
                if val_id:
                    migrated_ids.add(val_id)
        wb.close()

    df = pl.read_csv(funnel_csv_path, infer_schema_length=0, truncate_ragged_lines=True)
    cols = {c.strip().lower(): c for c in df.columns}

    id_col = cols.get('id', df.columns[0])
    status_col = cols.get('status', 'Status')
    reg_date_col = cols.get('registered date', 'Registered Date')
    onboard_date_col = cols.get('onboarded date', 'Onboarded Date')
    perm_addr_col = cols.get('permanent address', 'Permanent Address')

    raw_headers = df.columns
    base_headers = ['Ageing'] + list(raw_headers)
    base_rows = []
    distinct_dates = set()

    cutoff_dt = datetime.strptime(cutoff_date_str, '%Y-%m-%d').date()

    for row in df.iter_rows(named=True):
        emp_id = str(row.get(id_col) or '').strip().upper()
        if not emp_id or emp_id in migrated_ids:
            continue

        status = str(row.get(status_col) or '').strip()
        onboard_date = str(row.get(onboard_date_col) or '').strip()
        reg_date = str(row.get(reg_date_col) or '').strip()
        perm_addr = str(row.get(perm_addr_col) or '').strip().lower()

        if onboard_date and start_date_str <= onboard_date <= cutoff_date_str:
            if status.lower().startswith('onboarded'):
                is_hub_mig = 'hub migration' in perm_addr
                if not is_hub_mig:
                    distinct_dates.add(onboard_date)

                # Ageing Calculation
                ageing = 0
                if reg_date and reg_date != '-':
                    try:
                        reg_dt = datetime.strptime(reg_date[:10], '%Y-%m-%d').date()
                        diff = (cutoff_dt - reg_dt).days
                        ageing = max(0, diff)
                    except Exception:
                        ageing = 0

                row_vals = [ageing] + [row.get(c) for c in raw_headers]
                base_rows.append(row_vals)

    sorted_dates = sorted(list(distinct_dates))
    return sorted_dates, base_rows, base_headers
