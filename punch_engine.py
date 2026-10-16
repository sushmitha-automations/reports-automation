"""
Punch Adherence Engine.
Assembles total workforce headcount from Onroll HC & Valinor (LM, HL, FM) masters,
enriches supervisor roles from Chroma masters, joins daily biometric attendance punches,
and constructs Attendance Base, Punched Base, and Base data structures using Polars.
"""
import os
from typing import Dict, List, Tuple, Any, Optional
import polars as pl
import openpyxl

from config import (
    PUNCH_LM_PODS, PUNCH_HL_REGIONS, PUNCH_FM_REGIONS, PUNCH_HORO_SITES
)


def determine_vehicle_vertical(pod_name: str) -> str:
    """Classifies a POD/Site into LM, FM, HL, or HO-RO vertical."""
    p_upper = str(pod_name or '').strip().upper()
    if p_upper.startswith("FM_"):
        return "FM"
    if p_upper.startswith("HL_"):
        return "HL"
    if " - HO" in p_upper or " - RO" in p_upper or p_upper.endswith(" RO") or p_upper.endswith(" HO"):
        if "POD_" in p_upper:
            return "LM"
        return "HO-RO"
    return "LM"


def map_chroma_designation_to_role(desig: str) -> str:
    """Normalizes Chroma designations to punch report role abbreviations."""
    d = str(desig or '').strip().upper()
    if "TEAM LEADER" in d or d == "TL" or "FIELD TL" in d:
        return "TL"
    if "STORE MANAGER" in d:
        return "STORE MANAGER"
    if "PICKER" in d:
        return "PICKER"
    if "HELPER" in d:
        return "HELPER"
    if "DELIVERY" in d or "AMBASSADOR" in d or "DE" in d:
        return "DELIVERY AMBASSADORS"
    if "GC" in d:
        return "GC"
    if "FLEX" in d:
        return "Flex"
    if "IT EXECUTIVE" in d:
        return "IT EXECUTIVE"
    if "FINANCE" in d:
        return "FINANCE EXECUTIVE"
    if "MIS" in d:
        return "Mis Executive"
    if d:
        return str(desig).strip()
    return "Associate"


def normalize_lm_pod(pod: str, branch: str, fac_pod_map: Dict[str, str]) -> str:
    """Normalizes LM POD names, combining ROK/KL and NCR/UP_West."""
    p = str(pod or '').strip()
    b = str(branch or '').strip()
    if p in ('POD_ROK', 'POD_KL'):
        return 'POD_ROK/KL'
    if p in ('POD_NCR', 'POD_NCR/UP_West'):
        return 'POD_NCR/UP_West'
    if p in PUNCH_LM_PODS:
        return p

    f_p = fac_pod_map.get(b.lower(), '')
    if f_p in ('POD_ROK', 'POD_KL'):
        return 'POD_ROK/KL'
    if f_p in ('POD_NCR', 'POD_NCR/UP_West'):
        return 'POD_NCR/UP_West'
    if f_p in PUNCH_LM_PODS:
        return f_p

    b_up = b.upper()
    f_up = f_p.upper()
    if 'HO' in b_up or 'RO' in b_up or 'HO' in f_up or 'RO' in f_up:
        target = f_up if ('HO' in f_up or 'RO' in f_up) else b_up
        for s in PUNCH_HORO_SITES:
            if s.upper().replace(' ', '').replace('-', '') in target.replace(' ', '').replace('-', ''):
                return s
        return f_p if f_p else b
    return p if p else 'Not Applicable'


def normalize_hl_region(pod: str, branch: str, fac_pod_map: Dict[str, str]) -> str:
    """Normalizes HL PODs into 4 standard regions or HO-RO sites."""
    p = str(pod or '').strip()
    b = str(branch or '').strip()
    p_up = p.upper()

    if p in ('HL NCR', 'HL_NROI', 'HL_North') or 'NORTH' in p_up:
        return 'HL_North'
    if p in ('HL BOM/PNQ', 'HL_BOM/PNQ', 'HL_WROI', 'HL_West') or 'WEST' in p_up:
        return 'HL_West'
    if p in ('HL Bangalore', 'HL CHENNAI', 'HL HYDERABAD', 'HL Hyderabad', 'HL_South') or 'SOUTH' in p_up:
        return 'HL_South'
    if p in ('HL Kolkata', 'HL kolkata', 'HL_EROI', 'HL_East') or 'EAST' in p_up:
        return 'HL_East'

    f_p = fac_pod_map.get(b.lower(), '')
    if f_p:
        if f_p in PUNCH_HL_REGIONS:
            return f_p
        f_up = f_p.upper()
        if 'NORTH' in f_up:
            return 'HL_North'
        if 'WEST' in f_up:
            return 'HL_West'
        if 'SOUTH' in f_up:
            return 'HL_South'
        if 'EAST' in f_up:
            return 'HL_East'
        if 'HO' in f_up or 'RO' in f_up:
            for s in PUNCH_HORO_SITES:
                if s.upper().replace(' ', '').replace('-', '') in f_up.replace(' ', '').replace('-', ''):
                    return s
            return f_p

    b_up = b.upper()
    if 'HO' in b_up or 'RO' in b_up:
        for s in PUNCH_HORO_SITES:
            if s.upper().replace(' ', '').replace('-', '') in b_up.replace(' ', '').replace('-', ''):
                return s
        return b
    if 'BLR' in b_up:
        return 'HL_South'
    return p if p else 'Not Applicable'


def normalize_fm_region(pod: str, branch: str, fac_pod_map: Dict[str, str]) -> str:
    """Normalizes FM branch codes into 4 standard regions or HO-RO sites."""
    p = str(pod or '').strip()
    b = str(branch or '').strip()
    if p in PUNCH_FM_REGIONS:
        return p

    f_p = fac_pod_map.get(b.lower(), '')
    if f_p:
        if f_p in PUNCH_FM_REGIONS:
            return f_p
        f_up = f_p.upper()
        if 'NORTH' in f_up:
            return 'FM_North'
        if 'WEST' in f_up:
            return 'FM_West'
        if 'SOUTH' in f_up:
            return 'FM_South'
        if 'EAST' in f_up:
            return 'FM_East'
        if 'HO' in f_up or 'RO' in f_up:
            for s in PUNCH_HORO_SITES:
                if s.upper().replace(' ', '').replace('-', '') in f_up.replace(' ', '').replace('-', ''):
                    return s
            return f_p

    b_up = b.upper()
    if 'HO' in b_up or 'RO' in b_up:
        for s in PUNCH_HORO_SITES:
            if s.upper().replace(' ', '').replace('-', '') in b_up.replace(' ', '').replace('-', ''):
                return s
        return b
    if 'BLR' in b_up or 'BANGALORE' in b_up:
        return 'FM_South'
    if 'DEL' in b_up or 'NOI' in b_up or 'GGN' in b_up or 'ROHINI' in b_up or 'MUNDKA' in b_up:
        return 'FM_North'
    if 'ST_' in b_up or 'BOM' in b_up or 'AMD' in b_up or 'RJT' in b_up:
        return 'FM_West'
    if 'CCU' in b_up or 'HOWRAH' in b_up or 'PATNA' in b_up or 'PTN' in b_up:
        return 'FM_East'
    return 'FM_North'


def build_punch_adherence_dataset(
    onroll_path: str,
    offroll_path: Optional[str],
    attendance_csv_path: str,
    facility_master_csv_path: str,
    chroma_file_paths: List[str],
    date_list: List[str]
) -> Tuple[List[Dict[str, Any]], List[Tuple[str, str, str]]]:
    """
    Constructs the complete punch adherence dataset from Onroll HC and Valinor masters.
    Returns:
        workforce_records: list of dicts with employee attributes and daily punches
        base_hub_combos: list of unique (Hub Name, POD Name, Veticle) tuples for Base sheet
    """
    # 1. Load Facility Master POD Map
    df_fac = pl.read_csv(facility_master_csv_path, ignore_errors=True, infer_schema_length=0)
    fac_cols = {c.strip().lower(): c for c in df_fac.columns}
    fac_name_col = fac_cols.get('facility name') or df_fac.columns[1]
    pod_name_col = fac_cols.get('pod name') or df_fac.columns[5]

    fac_pod_map: Dict[str, str] = {}
    for r in df_fac.select([fac_name_col, pod_name_col]).iter_rows():
        f_name = str(r[0] or '').strip().lower()
        p_name = str(r[1] or '').strip()
        if f_name and p_name and p_name != 'None':
            fac_pod_map[f_name] = p_name

    # 2. Identify Valinor master files for LM, HL, and FM
    valinor_lm_path = None
    valinor_hl_path = None
    valinor_fm_path = None

    for c_path in chroma_file_paths:
        if not os.path.exists(c_path):
            continue
        base_low = os.path.basename(c_path).lower()
        if '(lm)' in base_low or 'valinor lm' in base_low or 'valinor (lm)' in base_low:
            valinor_lm_path = c_path
        elif '(fm)' in base_low or 'valinor fm' in base_low or 'valinor (fm)' in base_low:
            valinor_fm_path = c_path
        elif '(hl' in base_low or 'valinor hl' in base_low or 'valinor (hl' in base_low or 'corp' in base_low:
            valinor_hl_path = c_path

    # Extract cut-off day numbers for date indexing
    day_numbers = []
    for d_str in date_list:
        try:
            parts = d_str.split('-')
            day_numbers.append(int(parts[2]))
        except Exception:
            day_numbers.append(len(day_numbers) + 1)

    workforce: List[Dict[str, Any]] = []

    # 3. Load On-Roll Employees (Onroll HC.xlsx -> Main data)
    df_on = pl.read_excel(onroll_path, sheet_name='Main data')
    on_cols = {c.strip().lower(): c for c in df_on.columns}
    on_id_col = on_cols.get('employeecode') or df_on.columns[0]
    on_name_col = on_cols.get('employee name') or df_on.columns[1]
    on_sit_col = on_cols.get('sitting location (office location)') or df_on.columns[4]
    on_dept_col = on_cols.get('sub cost centre') or df_on.columns[6]
    on_pod_col = on_cols.get('pod') or df_on.columns[7]

    for r in df_on.iter_rows(named=True):
        emp_id = str(r.get(on_id_col) or '').strip()
        if not emp_id:
            continue
        emp_name = str(r.get(on_name_col) or '').strip()
        sit_loc = str(r.get(on_sit_col) or '').strip()
        dept = str(r.get(on_dept_col) or '').strip()
        raw_pod = str(r.get(on_pod_col) or '').strip()

        # Derive POD with Facility Master priority & normalization
        fac_p = fac_pod_map.get(sit_loc.lower(), '')
        if fac_p in ('POD_ROK', 'POD_KL'):
            derived_pod = 'POD_ROK/KL'
        elif fac_p in ('POD_NCR',):
            derived_pod = 'POD_NCR/UP_West'
        elif fac_p:
            derived_pod = fac_p
        elif 'HO' in sit_loc or 'RO' in sit_loc:
            derived_pod = sit_loc
        elif raw_pod in ('POD_ROK', 'POD_KL'):
            derived_pod = 'POD_ROK/KL'
        elif raw_pod in ('POD_NCR',):
            derived_pod = 'POD_NCR/UP_West'
        elif raw_pod and raw_pod not in ('Not Applicable', '-', 'None'):
            derived_pod = raw_pod
        else:
            derived_pod = 'Not Applicable'

        workforce.append({
            'id': emp_id,
            'name': emp_name,
            'type': 'On-Roll',
            'pod': derived_pod,
            'facility': sit_loc,
            'department': dept if dept else 'Central',
            'role': 'Employee',
            'punches': {}
        })

    # 4. Load Off-Roll Employees
    # Primary Source: Valinor (LM), Valinor (HL & Corp), Valinor (FM)
    if valinor_lm_path or valinor_hl_path or valinor_fm_path:
        valinor_specs = [
            (valinor_lm_path, 'LM', normalize_lm_pod),
            (valinor_hl_path, 'HL', normalize_hl_region),
            (valinor_fm_path, 'FM', normalize_fm_region)
        ]

        for v_path, v_name, norm_func in valinor_specs:
            if not v_path or not os.path.exists(v_path):
                continue
            try:
                wb = openpyxl.load_workbook(v_path, read_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                wb.close()
                if not rows:
                    continue

                headers = rows[0]
                status_c = 67
                pod_c = 23
                branch_c = 14
                desig_c = 12
                dept_c = 16

                # Dynamically locate headers if shifted
                for idx, h in enumerate(headers):
                    if h:
                        h_str = str(h).strip().lower()
                        if 'current status' in h_str:
                            status_c = idx
                        elif h_str in ('pod', 'pod name'):
                            pod_c = idx
                        elif 'branch code' in h_str:
                            branch_c = idx
                        elif 'designation' in h_str:
                            desig_c = idx
                        elif 'sub cost center' in h_str or 'sub cost centre' in h_str:
                            dept_c = idx

                for r in rows[1:]:
                    status_val = str(r[status_c] if status_c < len(r) else '').strip().lower()
                    if status_val != 'active':
                        continue
                    emp_id = str(r[0] if len(r) > 0 else '').strip()
                    if not emp_id or emp_id.lower() == 'none':
                        continue

                    emp_name = str(r[2] if len(r) > 2 else '').strip()
                    branch = str(r[branch_c] if branch_c < len(r) else '').strip()
                    desig = str(r[desig_c] if desig_c < len(r) else '').strip()
                    dept = str(r[dept_c] if dept_c < len(r) else '').strip() or f'{v_name} Operations'
                    raw_pod = str(r[pod_c] if pod_c < len(r) else '').strip()

                    pod = norm_func(raw_pod, branch, fac_pod_map)
                    role = map_chroma_designation_to_role(desig)

                    # Extract daily punches from columns 24..54 (col index = 23 + day_num)
                    punches = {}
                    for idx_d, d_str in enumerate(date_list):
                        day_n = day_numbers[idx_d]
                        col_idx = 23 + day_n
                        if col_idx < len(r):
                            cell_val = str(r[col_idx] or '').strip().upper()
                            punches[d_str] = 'P' if cell_val in ('P', 'NPP') else 'A'
                        else:
                            punches[d_str] = 'A'

                    workforce.append({
                        'id': emp_id,
                        'name': emp_name,
                        'type': 'Off-Roll',
                        'pod': pod,
                        'facility': branch,
                        'department': dept,
                        'role': role,
                        'punches': punches
                    })
            except Exception as e:
                print(f"[Warning] Error parsing Valinor file {v_path}: {e}")

    elif offroll_path and os.path.exists(offroll_path):
        # Fallback to Offroll HC.xlsx if Valinor files are absent
        df_off = pl.read_excel(offroll_path, sheet_name='Main Data')
        off_cols = {c.strip().lower(): c for c in df_off.columns}
        off_id_col = off_cols.get('empoyee id') or off_cols.get('employee id') or df_off.columns[0]
        off_name_col = off_cols.get('name') or df_off.columns[2]
        off_branch_col = off_cols.get('branch code') or df_off.columns[3]

        for r in df_off.iter_rows(named=True):
            emp_id = str(r.get(off_id_col) or '').strip()
            if not emp_id:
                continue
            emp_name = str(r.get(off_name_col) or '').strip()
            branch = str(r.get(off_branch_col) or '').strip()
            pod = fac_pod_map.get(branch.lower(), '') or 'Not Applicable'

            workforce.append({
                'id': emp_id,
                'name': emp_name,
                'type': 'Off-Roll',
                'pod': pod,
                'facility': branch,
                'department': 'Operations',
                'role': 'Associate',
                'punches': {}
            })

    # 5. Load Biometric Attendance Punches (Valinor Attendance.csv / .xlsx)
    if attendance_csv_path and os.path.exists(attendance_csv_path):
        if attendance_csv_path.lower().endswith(('.xlsx', '.xls')):
            try:
                df_att = pl.read_excel(attendance_csv_path)
                df_att = df_att.select([pl.col(c).cast(pl.String) for c in df_att.columns])
            except Exception:
                wb = openpyxl.load_workbook(attendance_csv_path, read_only=True, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                wb.close()
                if rows:
                    headers = [str(c or '').strip() for c in rows[0]]
                    data_dict = {h: [str(r[i] or '').strip() if i < len(r) else '' for r in rows[1:]] for i, h in enumerate(headers)}
                    df_att = pl.DataFrame(data_dict)
                else:
                    df_att = pl.DataFrame()
        else:
            df_att = pl.read_csv(attendance_csv_path, ignore_errors=True, infer_schema_length=0)

        att_cols = {str(c).strip().lower(): c for c in df_att.columns}
        att_id_col = att_cols.get('employeeid') or df_att.columns[1]

        # Map dates in attendance file: '01-10-2026' or '2026-10-01' -> date_list keys
        att_date_cols: Dict[str, str] = {}
        for col in df_att.columns:
            clean_c = str(col).strip()
            parts = clean_c.split('-')
            # DD-MM-YYYY
            if len(parts) == 3 and len(parts[0]) == 2 and parts[2][:4].isdigit():
                try:
                    d_day = parts[0]
                    m_num = int(parts[1])
                    y_val = parts[2][:4]
                    y_short = y_val[-2:]
                    m_names = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
                    m_str = m_names[m_num - 1]
                    header_fmt = f"{d_day}-{m_str}-{y_short}"
                    iso_fmt = f"{int(y_val):04d}-{m_num:02d}-{int(d_day):02d}"
                    att_date_cols[header_fmt] = clean_c
                    att_date_cols[iso_fmt] = clean_c
                except Exception:
                    pass
            # YYYY-MM-DD
            elif len(parts) == 3 and len(parts[0]) == 4:
                try:
                    y_val = parts[0]
                    m_num = int(parts[1])
                    d_day = parts[2][:2]
                    y_short = y_val[-2:]
                    m_names = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
                    m_str = m_names[m_num - 1]
                    header_fmt = f"{d_day}-{m_str}-{y_short}"
                    iso_fmt = f"{int(y_val):04d}-{m_num:02d}-{int(d_day):02d}"
                    att_date_cols[header_fmt] = clean_c
                    att_date_cols[iso_fmt] = clean_c
                except Exception:
                    pass

        att_punch_map: Dict[str, Dict[str, str]] = {}
        for r in df_att.iter_rows(named=True):
            raw_emp_id = str(r.get(att_id_col) or '').strip()
            if not raw_emp_id:
                continue
            p_dict = {}
            for d_hdr in date_list:
                orig_col = att_date_cols.get(d_hdr) or (d_hdr if d_hdr in r else None)
                if orig_col and orig_col in r:
                    val = str(r.get(orig_col) or '').strip().upper()
                    p_dict[d_hdr] = 'P' if val == 'P' else 'A'
                else:
                    p_dict[d_hdr] = 'A'
            att_punch_map[raw_emp_id] = p_dict

        # Merge punches into workforce records
        for emp in workforce:
            emp_id = emp['id']
            clean_numeric_id = emp_id.replace('SF', '').replace('PROD', '').strip()
            p_data = att_punch_map.get(emp_id) or att_punch_map.get(clean_numeric_id) or {}

            # If employee already has punches (from Valinor files), union with Attendance CSV
            if emp['punches']:
                for d_hdr in date_list:
                    if p_data.get(d_hdr) == 'P':
                        emp['punches'][d_hdr] = 'P'
            else:
                for d_hdr in date_list:
                    emp['punches'][d_hdr] = p_data.get(d_hdr, 'A')

    # 6. Extract unique Hub & POD combinations for Base sheet
    combo_set = set()
    base_hub_combos: List[Tuple[str, str, str]] = []
    for emp in workforce:
        hub = emp['facility']
        pod = emp['pod']
        key = (hub, pod)
        if key not in combo_set:
            combo_set.add(key)
            vehicle = determine_vehicle_vertical(pod)
            base_hub_combos.append((hub, pod, vehicle))

    return workforce, base_hub_combos
