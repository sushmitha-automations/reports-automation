"""
Excel Report Builder using OpenPyXL.
Applies exact cell styling, fonts, borders, fills, number formats, column dimensions,
and live Excel formulas strictly matching reference reports with zero over-styling.
"""
import os
from typing import Dict, List, Any, Optional, Callable, Tuple
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

from config import (
    LM_PODS, HL_REGIONS, FM_REGIONS, HORO_SITES,
    PUNCH_LM_PODS, PUNCH_HL_REGIONS, PUNCH_FM_REGIONS, PUNCH_HORO_SITES,
    REF_VAL_RULES, REF_ROLE_RULES,
    BORDER_THIN_GRAY, BORDER_THIN_BLACK, BORDER_DOUBLE_BOTTOM,
    FILL_DARK_BLUE, FILL_LIGHT_BLUE, FILL_SOFT_BLUE_HDR,
    FILL_HEADER_BLUE_LIGHT_80, FILL_GRAND_TOTAL_DARK_25,
    FILL_YELLOW, FILL_RED, FILL_LIGHT_GRAY,
    FILL_PUNCH_HDR, FILL_PUNCH_RED, FILL_PUNCH_AMBER, FILL_PUNCH_GREEN, FILL_PUNCH_TOTAL,
    FONT_CALIBRI_11, FONT_CALIBRI_11_BOLD, FONT_CALIBRI_11_BOLD_WHITE,
    FONT_APTOS_11, FONT_APTOS_11_BOLD, FONT_APTOS_11_BOLD_WHITE, FONT_APTOS_9_BOLD,
    ALIGN_CENTER, ALIGN_CENTER_WRAP, ALIGN_LEFT,
    FMT_TEXT, FMT_PERCENT_2DEC, FMT_PERCENT_INT, FMT_INT_COMMAS
)

def format_date_header(dt_str: str) -> str:
    """Formats '2026-09-01' into '01-Sep-26'."""
    try:
        parts = dt_str.split('-')
        if len(parts) == 3:
            months = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
            m_idx = int(parts[1]) - 1
            m_str = months[m_idx] if 0 <= m_idx < 12 else parts[1]
            yr = parts[0][-2:]
            return f"{parts[2]}-{m_str}-{yr}"
    except Exception:
        pass
    return dt_str

# ==============================================================================
# ATTENDANCE ADHERENCE REPORT GENERATOR
# ==============================================================================

def build_attendance_adherence_report(
    output_path: str,
    date_list: List[str],
    evaluated_val_base: List[Dict[str, Any]],
    user_rows: List[Dict[str, Any]],
    facilities: List[Dict[str, Any]],
    val_att_rows: List[Dict[str, Any]],
    chroma_rows: List[Dict[str, Any]],
    val_dump_role_map: Dict[int, str],
    progress_callback: Optional[Callable[[float, str], None]] = None
):
    """
    Constructs the 11-sheet Attendance Adherence Report with exact Calibri 11pt formatting,
    live COUNTIFS formulas, and strict @ text formatting for IDs/Aadhaar/Mobile.
    """
    def notify(pct: float, msg: str):
        if progress_callback:
            progress_callback(pct, msg)

    notify(0.05, "Creating worksheets & writing Facility Master...")
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Instantiate worksheets in frontmost order
    ws_lm = wb.create_sheet(title='Attendance Check - LM')
    ws_hl = wb.create_sheet(title='Attendance Check - HL')
    ws_fm = wb.create_sheet(title='Attendance Check - FM')
    ws_val = wb.create_sheet(title='Validations Base')
    ws_chroma_val_sum = wb.create_sheet(title='Chroma vs Valinor summary')
    ws_chroma = wb.create_sheet(title='Chroma Dump')
    ws_val_att = wb.create_sheet(title='Valinor Attendance')
    ws_ref_role = wb.create_sheet(title='Reference - Role')
    ws_ref_val = wb.create_sheet(title='Ref - Validation')
    ws_val_dump = wb.create_sheet(title='Valinor Dump')
    ws_fac = wb.create_sheet(title='Valinor Dump - Facility')

    num_days = len(date_list)

    # -------------------------------------------------------------------------
    # SHEET 11: Valinor Dump - Facility
    # -------------------------------------------------------------------------
    fac_headers = ['Facility id', 'Facility Name', 'type', 'saruman id', 'darsa id', 'Pod Name', 'Pod Id', 'State Name', 'State Id']
    ws_fac.append(fac_headers)
    for cell in ws_fac[1]:
        cell.font = FONT_CALIBRI_11_BOLD_WHITE
        cell.fill = FILL_DARK_BLUE
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY

    for f in facilities:
        ws_fac.append([
            int(f['id']) if str(f['id']).isdigit() else f['id'],
            f['name'], f['type'],
            int(f['saruman_id']) if str(f['saruman_id']).isdigit() else f['saruman_id'],
            int(f['darsa_id']) if str(f['darsa_id']).isdigit() else f['darsa_id'],
            f['pod_name'],
            int(f['pod_id']) if str(f['pod_id']).isdigit() else f['pod_id'],
            f['state_name'],
            int(f['state_id']) if str(f['state_id']).isdigit() else f['state_id']
        ])
    for row in ws_fac.iter_rows(min_row=2):
        for cell in row:
            cell.font = FONT_CALIBRI_11
            cell.border = BORDER_THIN_GRAY

    # -------------------------------------------------------------------------
    # SHEET 10: Valinor Dump
    # -------------------------------------------------------------------------
    vd_headers = ['id', 'SFX ID', 'phone_number', 'name', 'status', 'aadhaar_number', 'registered_facility', 'created_date_ist', 'role_type', 'rolename', 'facility_name', 'POD']
    ws_val_dump.append(vd_headers)
    for cell in ws_val_dump[1]:
        cell.font = FONT_CALIBRI_11_BOLD_WHITE
        cell.fill = FILL_DARK_BLUE
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY

    notify(0.12, f"Writing Valinor User Dump ({len(user_rows):,} records)...")
    total_u = len(user_rows)
    for r_idx, u in enumerate(user_rows, start=2):
        if total_u > 0 and r_idx % 20000 == 0:
            notify(0.12 + (r_idx / total_u) * 0.38, f"Writing Valinor Dump: {r_idx:,} / {total_u:,} rows...")
        row_cells = [
            u['numericId'],
            f"=CONCATENATE(\"SFXPROD\",A{r_idx})",
            u['phone_number'],
            u['name'],
            u['status'],
            u['aadhaar_number'],
            u['registered_facility'],
            u['created_date_ist'],
            u['role_type'],
            u['rolename'],
            u['facility_name'],
            u['pod']
        ]
        ws_val_dump.append(row_cells)
        # Explicit text format for phone and aadhaar
        ws_val_dump.cell(row=r_idx, column=3).number_format = FMT_TEXT
        ws_val_dump.cell(row=r_idx, column=6).number_format = FMT_TEXT
        for c_idx in range(1, 13):
            cell = ws_val_dump.cell(row=r_idx, column=c_idx)
            cell.font = FONT_CALIBRI_11
            cell.border = BORDER_THIN_GRAY

    # -------------------------------------------------------------------------
    # SHEET 9: Ref - Validation
    # -------------------------------------------------------------------------
    notify(0.50, "Writing Reference Rules (Validation & Role)...")
    ws_ref_val.append(['Chroma Status', 'Valinor Status', 'Validation Status'])
    for cell in ws_ref_val[1]:
        cell.font = FONT_CALIBRI_11_BOLD_WHITE
        cell.fill = FILL_DARK_BLUE
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY
    for rule in REF_VAL_RULES:
        ws_ref_val.append(list(rule))
    for row in ws_ref_val.iter_rows(min_row=2):
        for cell in row:
            cell.font = FONT_CALIBRI_11
            cell.border = BORDER_THIN_GRAY

    # -------------------------------------------------------------------------
    # SHEET 8: Reference - Role
    # -------------------------------------------------------------------------
    ws_ref_role.append(['Valinor Role Name', 'Chroma Designation'])
    for cell in ws_ref_role[1]:
        cell.font = FONT_CALIBRI_11_BOLD_WHITE
        cell.fill = FILL_DARK_BLUE
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY
    for rule in REF_ROLE_RULES:
        ws_ref_role.append(list(rule))
    for row in ws_ref_role.iter_rows(min_row=2):
        for cell in row:
            cell.font = FONT_CALIBRI_11
            cell.border = BORDER_THIN_GRAY

    # -------------------------------------------------------------------------
    # SHEET 7: Valinor Attendance
    # -------------------------------------------------------------------------
    notify(0.55, f"Writing Valinor Biometric Attendance punches ({len(val_att_rows):,} rows)...")
    va_headers = ['ID', 'user_id', 'user_name', 'role'] + date_list
    ws_val_att.append(va_headers)
    for cell in ws_val_att[1]:
        cell.font = FONT_CALIBRI_11_BOLD_WHITE
        cell.fill = FILL_DARK_BLUE
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY

    for r_idx, rec in enumerate(val_att_rows, start=2):
        num_id = rec['empId']
        sfx_id = f"SFXPROD{num_id}" if num_id else ""
        role_val = val_dump_role_map.get(num_id, 'DE')
        row_data = [num_id, sfx_id, rec['name'], role_val]
        for d in date_list:
            row_data.append(rec['days'].get(d, 'A'))
        ws_val_att.append(row_data)
        for c_idx in range(1, len(va_headers) + 1):
            cell = ws_val_att.cell(row=r_idx, column=c_idx)
            cell.font = FONT_CALIBRI_11
            cell.border = BORDER_THIN_GRAY

    # -------------------------------------------------------------------------
    # SHEET 6: Chroma Dump
    # -------------------------------------------------------------------------
    notify(0.68, f"Writing Chroma Dump ({len(evaluated_val_base):,} rows)...")
    c_headers = [
        'Mobile Check', 'Aadhar Check', 'Valinor ID', 'Emp ID', 'App ID', 'NAME', 'DOJ', 'DOB',
        'AGE', "FATHER'S NAME", 'MOBILE', 'AADHAR NO.', 'PAN Number / ID Number', 'GENDER',
        'SKILL TYPE', 'DESIGNATION', 'ROLE', 'Payroll Source', 'DEPARTMENT', 'BRANCH CODE',
        'DARSA BRANCH CODE', 'COST CENTRE', 'COST CENTRE SUB UNIT', 'STATE', 'CITY', 'Vendor Name',
        'R.M Employee ID', 'R.M Name', 'STATUS (R.M.)', 'POD Name'
    ] + [str(d) for d in range(1, num_days + 1)] + [
        'P', 'PL', 'HPL', 'WO', 'HD', 'A', 'LWP', 'HL', 'NA', 'Total', 'Payable Days',
        'LWP Counter', 'Remarks', 'Total Leave availed', 'CURRENT STATUS', 'LWD',
        'HRBP Name', 'HRBP Employee ID', 'Bad Apple', 'Exit Reason'
    ]
    ws_chroma.append(c_headers)
    for cell in ws_chroma[1]:
        cell.font = FONT_CALIBRI_11_BOLD_WHITE
        cell.fill = FILL_DARK_BLUE
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY

    for r_idx, eval_item in enumerate(evaluated_val_base, start=2):
        c_item = eval_item['chroma_item']
        row_cells = [
            f"=XLOOKUP(K{r_idx},'Valinor Dump'!C:C,'Valinor Dump'!A:A,\"\")",
            f"=XLOOKUP(L{r_idx},'Valinor Dump'!F:F,'Valinor Dump'!A:A,\"\")",
            eval_item['valinor_id'],
            c_item['emp_id'],
            "",  # App ID
            c_item['name'],
            c_item['doj'],
            c_item['dob'],
            c_item['age'],
            c_item['father_name'],
            c_item['mobile'],
            c_item['aadhaar'],
            c_item['pan'],
            c_item['gender'],
            c_item['skill'],
            c_item['designation'],
            "",  # ROLE
            c_item['payroll_source'],
            "Operations",
            c_item['branch_code'],
            c_item['branch_code'],
            c_item['cost_center'],
            c_item['sub_cost_center'],
            c_item['state'],
            c_item['city'],
            c_item['vendor_name'],
            c_item['rm_emp_id'],
            c_item['rm_name'],
            c_item['rm_status'],
            eval_item['chroma_pod']
        ] + c_item['date_punches']

        # Append trailing stats from raw Chroma if available
        raw_r = c_item['raw_row']
        trailing_start = 53
        for col_t in range(trailing_start, len(raw_r)):
            row_cells.append(raw_r[col_t])

        ws_chroma.append(row_cells)
        # Formatting mobile and aadhaar as text
        ws_chroma.cell(row=r_idx, column=11).number_format = FMT_TEXT
        ws_chroma.cell(row=r_idx, column=12).number_format = FMT_TEXT
        for c_idx in range(1, len(row_cells) + 1):
            cell = ws_chroma.cell(row=r_idx, column=c_idx)
            cell.font = FONT_CALIBRI_11
            cell.border = BORDER_THIN_GRAY

    # -------------------------------------------------------------------------
    # SHEET 4: Validations Base
    # -------------------------------------------------------------------------
    notify(0.78, "Writing Validations Base & Chroma vs Valinor summary...")
    v_h1 = [''] * 16 + ['Overall'] * 7
    for d in date_list:
        v_h1.extend([d, d, d])
    ws_val.append(v_h1)

    v_h2 = [
        'Valinor ID', 'Valinor DOJ', 'Chroma ID', 'Chroma NAME', 'DOJ',
        'Valinor Role', 'DESIGNATION', 'Valinor Facility', 'Valinor POD', 'Chroma Location',
        'POD', 'Department', 'Role Issue', 'Sitting Location Issue',
        'DOJ Gap for Aug\'26 Joiners', 'DOJ Gap Type',
        'MATCH', 'P in Chroma, A in Valinor', 'WO/PL/HD in Chroma, A in Valinor',
        'Tech Issue', 'NA in Chroma, P in Valinor', 'A/LWP in Chroma, P in Valinor',
        'WO/PL/HD in Chroma, P in Valinor'
    ]
    for d in date_list:
        v_h2.extend(['Chroma', 'Valinor', 'Validation'])
    ws_val.append(v_h2)

    # Style Header rows
    for r_num in (1, 2):
        for cell in ws_val[r_num]:
            cell.font = FONT_CALIBRI_11_BOLD_WHITE
            cell.fill = FILL_DARK_BLUE
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN_GRAY

    # Merge header cells
    ws_val.merge_cells('Q1:W1')
    for d_idx in range(num_days):
        sc = 24 + d_idx * 3
        ec = sc + 2
        ws_val.merge_cells(start_row=1, start_column=sc, end_row=1, end_column=ec)

    for r_idx, eval_item in enumerate(evaluated_val_base, start=3):
        c_item = eval_item['chroma_item']
        v_user = eval_item['valinor_user']

        val_role = v_user['rolename'] if v_user else 'DE'
        row_data = [
            eval_item['valinor_id'],
            v_user['created_date_ist'] if v_user else '',
            c_item['emp_id'],
            c_item['name'],
            c_item['doj'],
            val_role,
            c_item['designation'],
            v_user['registered_facility'] if v_user else '',
            eval_item['valinor_pod'],
            c_item['branch_code'],
            eval_item['chroma_pod'],
            "Operations",
            "N", "N", "", "",  # Issues
            eval_item['match_count'],
            eval_item['p_chroma_a_val'],
            eval_item['wo_chroma_a_val'],
            eval_item['tech_issue'],
            eval_item['na_chroma_p_val'],
            eval_item['a_chroma_p_val'],
            eval_item['wo_chroma_p_val']
        ]

        for d_idx in range(num_days):
            row_data.append(eval_item['date_punches'][d_idx])
            row_data.append(eval_item['date_val_att'][d_idx])
            row_data.append(eval_item['date_validations'][d_idx])

        ws_val.append(row_data)
        for c_idx in range(1, len(row_data) + 1):
            cell = ws_val.cell(row=r_idx, column=c_idx)
            cell.font = FONT_CALIBRI_11
            cell.border = BORDER_THIN_GRAY

    # -------------------------------------------------------------------------
    # SHEET 5: Chroma vs Valinor summary
    # -------------------------------------------------------------------------
    ws_chroma_val_sum.append([''] * 8)
    ws_chroma_val_sum.append(['CURRENT STATUS', 'Active', '', '', 'HRBP Name', 'Not Onboarded', 'Onboarded', 'Grand Total'])
    for c_idx in (1, 2, 5, 6, 7, 8):
        cell = ws_chroma_val_sum.cell(row=2, column=c_idx)
        cell.font = FONT_CALIBRI_11_BOLD_WHITE
        cell.fill = FILL_DARK_BLUE
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY

    # Group by HRBP
    hrbp_map: Dict[str, Dict[str, int]] = {}
    for eval_item in evaluated_val_base:
        raw_r = eval_item['chroma_item']['raw_row']
        hrbp = str(raw_r[71] if len(raw_r) > 71 and raw_r[71] else 'Unassigned').strip()
        is_onboarded = bool(eval_item['valinor_id'])
        if hrbp not in hrbp_map:
            hrbp_map[hrbp] = {'onboarded': 0, 'not_onboarded': 0}
        if is_onboarded:
            hrbp_map[hrbp]['onboarded'] += 1
        else:
            hrbp_map[hrbp]['not_onboarded'] += 1

    r_sum_idx = 3
    for hrbp, counts in sorted(hrbp_map.items()):
        ws_chroma_val_sum.append([
            "", "", "", "",
            hrbp, counts['not_onboarded'], counts['onboarded'],
            f"=F{r_sum_idx}+G{r_sum_idx}"
        ])
        for c_idx in range(5, 9):
            cell = ws_chroma_val_sum.cell(row=r_sum_idx, column=c_idx)
            cell.font = FONT_CALIBRI_11
            cell.border = BORDER_THIN_GRAY
            cell.alignment = ALIGN_LEFT if c_idx == 5 else ALIGN_CENTER
        r_sum_idx += 1

    # -------------------------------------------------------------------------
    # SHEETS 1, 2, 3: Attendance Check (LM, HL, FM)
    # -------------------------------------------------------------------------
    def populate_attendance_check_sheet(ws, label_header, entities, match_col_letter):
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 30

        r1 = [label_header]
        r2 = [label_header]
        for d in date_list:
            formatted_d = format_date_header(d)
            r1.extend([formatted_d, formatted_d, formatted_d, formatted_d])
            r2.extend(['Total reported in chroma', 'Matches with valinor', 'Not Matching', 'Adherence %'])

        ws.append(r1)
        ws.append(r2)

        # Style header rows BEFORE merging to ensure all cells retain border and fill
        for r_num in (1, 2):
            for cell in ws[r_num]:
                cell.font = FONT_CALIBRI_11_BOLD
                cell.fill = FILL_HEADER_BLUE_LIGHT_80
                cell.alignment = ALIGN_CENTER_WRAP
                cell.border = BORDER_THIN_GRAY

        ws.merge_cells('A1:A2')
        for d_idx in range(num_days):
            sc = 2 + d_idx * 4
            ec = sc + 3
            ws.merge_cells(start_row=1, start_column=sc, end_row=1, end_column=ec)

        # Data rows
        start_row = 3
        for e_idx, entity in enumerate(entities):
            cur_r = start_row + e_idx
            ws.row_dimensions[cur_r].height = 19
            row_data = [entity]

            for d_idx in range(num_days):
                # Validations Base column offsets
                # Date punch col = 24 + d_idx * 3
                # Validation col = 26 + d_idx * 3
                punch_col_let = get_column_letter(24 + d_idx * 3)
                val_col_let = get_column_letter(26 + d_idx * 3)

                tot_f = f"=COUNTIFS('Validations Base'!${match_col_letter}:${match_col_letter}, $A{cur_r}, 'Validations Base'!${punch_col_let}:${punch_col_let}, \"P\") + COUNTIFS('Validations Base'!${match_col_letter}:${match_col_letter}, $A{cur_r}, 'Validations Base'!${punch_col_let}:${punch_col_let}, \"NPP\")"
                
                tot_col_let = get_column_letter(2 + d_idx * 4)
                match_col_cur = get_column_letter(3 + d_idx * 4)
                
                match_f = f"=COUNTIFS('Validations Base'!${match_col_letter}:${match_col_letter}, $A{cur_r}, 'Validations Base'!${val_col_let}:${val_col_let}, \"MATCH\", 'Validations Base'!${punch_col_let}:${punch_col_let}, \"P\") + COUNTIFS('Validations Base'!${match_col_letter}:${match_col_letter}, $A{cur_r}, 'Validations Base'!${val_col_let}:${val_col_let}, \"MATCH\", 'Validations Base'!${punch_col_let}:${punch_col_let}, \"NPP\")"
                not_match_f = f"={tot_col_let}{cur_r}-{match_col_cur}{cur_r}"
                adh_f = f"=IF({tot_col_let}{cur_r}>0,{match_col_cur}{cur_r}/{tot_col_let}{cur_r},1)"

                row_data.extend([tot_f, match_f, not_match_f, adh_f])

            ws.append(row_data)

            # Apply cell formats
            for c_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=cur_r, column=c_idx)
                cell.font = FONT_CALIBRI_11
                cell.border = BORDER_THIN_GRAY
                if c_idx == 1:
                    cell.alignment = ALIGN_LEFT
                else:
                    cell.alignment = ALIGN_CENTER
                    if (c_idx - 2) % 4 == 3:
                        cell.number_format = FMT_PERCENT_INT
                    else:
                        cell.number_format = FMT_INT_COMMAS

        # Grand Total Row
        tot_row_num = start_row + len(entities)
        ws.row_dimensions[tot_row_num].height = 21
        tot_row_data = ['Grand Total']

        for d_idx in range(num_days):
            c_tot_let = get_column_letter(2 + d_idx * 4)
            c_match_let = get_column_letter(3 + d_idx * 4)
            c_not_let = get_column_letter(4 + d_idx * 4)

            tot_sum_f = f"=SUM({c_tot_let}3:{c_tot_let}{tot_row_num - 1})"
            match_sum_f = f"=SUM({c_match_let}3:{c_match_let}{tot_row_num - 1})"
            not_sum_f = f"=SUM({c_not_let}3:{c_not_let}{tot_row_num - 1})"
            adh_sum_f = f"=IF({c_tot_let}{tot_row_num}>0,{c_match_let}{tot_row_num}/{c_tot_let}{tot_row_num},1)"

            tot_row_data.extend([tot_sum_f, match_sum_f, not_sum_f, adh_sum_f])

        ws.append(tot_row_data)
        for c_idx in range(1, len(tot_row_data) + 1):
            cell = ws.cell(row=tot_row_num, column=c_idx)
            cell.font = FONT_CALIBRI_11_BOLD
            cell.fill = FILL_GRAND_TOTAL_DARK_25
            cell.border = BORDER_THIN_GRAY
            cell.alignment = ALIGN_LEFT if c_idx == 1 else ALIGN_CENTER
            if c_idx > 1 and (c_idx - 2) % 4 == 3:
                cell.number_format = FMT_PERCENT_INT
            elif c_idx > 1:
                cell.number_format = FMT_INT_COMMAS

        # Adherence % - Green Yellow - Red Color Scale (Min: Red F8696B, 50th: Yellow FFEB84, Max: Green 63BE7B)
        adh_ranges = []
        for d_idx in range(num_days):
            adh_col_let = get_column_letter(5 + d_idx * 4)
            adh_ranges.append(f"{adh_col_let}3:{adh_col_let}{tot_row_num - 1}")
        if adh_ranges:
            combined_adh_range = " ".join(adh_ranges)
            color_scale_rule = ColorScaleRule(
                start_type='min', start_color='F8696B',
                mid_type='percentile', mid_value=50, mid_color='FFEB84',
                end_type='max', end_color='63BE7B'
            )
            ws.conditional_formatting.add(combined_adh_range, color_scale_rule)

        # Column widths
        ws.column_dimensions['A'].width = 22 if label_header == 'POD' else 14
        for c_idx in range(2, len(tot_row_data) + 1):
            c_let = get_column_letter(c_idx)
            mod = (c_idx - 2) % 4
            if mod == 0:
                ws.column_dimensions[c_let].width = 23.3
            elif mod == 1:
                ws.column_dimensions[c_let].width = 19.9
            elif mod == 2:
                ws.column_dimensions[c_let].width = 13.1
            else:
                ws.column_dimensions[c_let].width = 12.7

    notify(0.88, "Writing Executive Summary sheets (Attendance Check LM, HL, FM)...")
    populate_attendance_check_sheet(ws_lm, 'POD', LM_PODS, 'K')
    populate_attendance_check_sheet(ws_hl, 'Region', HL_REGIONS, 'I')
    populate_attendance_check_sheet(ws_fm, 'Region', FM_REGIONS, 'I')

    notify(0.96, "Saving Attendance Adherence workbook to disk...")
    wb.save(output_path)
    wb.close()
    notify(1.00, "Attendance Adherence Workbook generated.")


# ==============================================================================
# ONBOARDING REPORT GENERATOR
# ==============================================================================

def build_onboarding_report(
    output_path: str,
    date_list: List[str],
    base_rows: List[List[Any]],
    base_headers: List[str],
    progress_callback: Optional[Callable[[float, str], None]] = None
):
    """
    Constructs the 9-sheet Onboarding Report with exact Aptos Narrow 11pt formatting,
    Aadhaar compliance tables, MTD summaries, and HR-wise manual approvals.
    """
    def notify(pct: float, msg: str):
        if progress_callback:
            progress_callback(pct, msg)

    notify(0.05, f"Creating worksheets & writing Onboarding Base ({len(base_rows):,} rows)...")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_lm = wb.create_sheet(title='Aadhar - LM')
    ws_hl = wb.create_sheet(title='Aadhar - HL')
    ws_fm = wb.create_sheet(title='Aadhar - FM')
    ws_horo = wb.create_sheet(title='Aadhar - HO-RO')
    ws_man_lm = wb.create_sheet(title='Manual - LM')
    ws_man_hl = wb.create_sheet(title='Manual - HL')
    ws_man_fm = wb.create_sheet(title='Manual - FM')
    ws_man_horo = wb.create_sheet(title='Manual - HO-RO')
    ws_base = wb.create_sheet(title='Onboarding Base')

    num_days = len(date_list)

    # -------------------------------------------------------------------------
    # SHEET 9: Onboarding Base
    # -------------------------------------------------------------------------
    ws_base.append(base_headers)
    for cell in ws_base[1]:
        cell.font = FONT_APTOS_11_BOLD
        cell.fill = FILL_HEADER_BLUE_LIGHT_80
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY

    for row_vals in base_rows:
        ws_base.append(row_vals)
    for row in ws_base.iter_rows(min_row=2):
        for cell in row:
            cell.font = FONT_APTOS_11
            cell.border = BORDER_THIN_GRAY

    # -------------------------------------------------------------------------
    # SHEETS 1, 2, 3, 4: Aadhar Summary Sheets
    # -------------------------------------------------------------------------
    def populate_aadhar_summary(ws, label_header, pod_list):
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 20

        r1 = [label_header]
        r2 = [label_header]

        for d in date_list:
            formatted_d = format_date_header(d)
            r1.extend([formatted_d, formatted_d, formatted_d, formatted_d, 'Hyperverge Red'])
            r2.extend(['Total Onboarding', 'Aadhar', 'Manual', 'Compliance %', 'Red'])

        ws.append(r1)
        ws.append(r2)

        # Style Headers BEFORE merging to ensure all cells get border and fill
        for c_idx in range(1, len(r1) + 1):
            cell_1 = ws.cell(row=1, column=c_idx)
            cell_2 = ws.cell(row=2, column=c_idx)

            cell_1.border = BORDER_THIN_GRAY
            cell_2.border = BORDER_THIN_GRAY
            cell_1.alignment = ALIGN_CENTER
            cell_2.alignment = ALIGN_CENTER

            if c_idx == 1:
                cell_1.font = FONT_APTOS_11_BOLD
                cell_1.fill = FILL_HEADER_BLUE_LIGHT_80
                cell_2.font = FONT_APTOS_11_BOLD
                cell_2.fill = FILL_HEADER_BLUE_LIGHT_80
            else:
                mod = (c_idx - 2) % 5
                if mod == 4:
                    cell_1.font = FONT_APTOS_11_BOLD_WHITE
                    cell_1.fill = FILL_RED
                    cell_2.font = FONT_APTOS_11_BOLD_WHITE
                    cell_2.fill = FILL_RED
                else:
                    cell_1.font = FONT_APTOS_11_BOLD
                    cell_1.fill = FILL_HEADER_BLUE_LIGHT_80
                    if mod == 3:
                        cell_2.font = FONT_APTOS_11_BOLD
                        cell_2.fill = FILL_YELLOW
                    else:
                        cell_2.font = FONT_APTOS_11_BOLD
                        cell_2.fill = FILL_HEADER_BLUE_LIGHT_80

        # Merge Header cells
        ws.merge_cells('A1:A2')
        for d_idx in range(num_days):
            sc = 2 + d_idx * 5
            ec = sc + 3
            ws.merge_cells(start_row=1, start_column=sc, end_row=1, end_column=ec)

        # Daily Table Data Rows
        start_row = 3
        for p_idx, pod in enumerate(pod_list):
            cur_r = start_row + p_idx
            ws.row_dimensions[cur_r].height = 19
            row_data = [pod]

            for d_idx in range(num_days):
                d_hdr_let = get_column_letter(2 + d_idx * 5)
                tot_let = get_column_letter(2 + d_idx * 5)
                aadh_let = get_column_letter(3 + d_idx * 5)

                f_tot = f"=COUNTIFS('Onboarding Base'!$K:$K, $A{cur_r}, 'Onboarding Base'!$M:$M, {d_hdr_let}$1, 'Onboarding Base'!$D:$D, \"Onboarded*\", 'Onboarding Base'!$AE:$AE, \"<>*hub migration*\")"
                f_aadh = f"=COUNTIFS('Onboarding Base'!$K:$K, $A{cur_r}, 'Onboarding Base'!$M:$M, {d_hdr_let}$1, 'Onboarding Base'!$T:$T, \"Y\", 'Onboarding Base'!$D:$D, \"Onboarded*\", 'Onboarding Base'!$AE:$AE, \"<>*hub migration*\")"
                f_man = f"=COUNTIFS('Onboarding Base'!$K:$K, $A{cur_r}, 'Onboarding Base'!$M:$M, {d_hdr_let}$1, 'Onboarding Base'!$T:$T, \"N\", 'Onboarding Base'!$D:$D, \"Onboarded*\", 'Onboarding Base'!$AE:$AE, \"<>*hub migration*\")"
                f_comp = f"=IFERROR({aadh_let}{cur_r}/{tot_let}{cur_r}, 1)"
                f_red = f"=COUNTIFS('Onboarding Base'!$K:$K, $A{cur_r}, 'Onboarding Base'!$M:$M, {d_hdr_let}$1, 'Onboarding Base'!$X:$X, \"red\", 'Onboarding Base'!$D:$D, \"Onboarded*\", 'Onboarding Base'!$AE:$AE, \"<>*hub migration*\")"

                row_data.extend([f_tot, f_aadh, f_man, f_comp, f_red])

            ws.append(row_data)

            for c_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=cur_r, column=c_idx)
                cell.font = FONT_APTOS_11
                cell.border = BORDER_THIN_GRAY
                if c_idx == 1:
                    cell.alignment = ALIGN_LEFT
                else:
                    cell.alignment = ALIGN_CENTER
                    mod = (c_idx - 2) % 5
                    if mod == 3:
                        cell.number_format = FMT_PERCENT_INT
                    else:
                        cell.number_format = FMT_INT_COMMAS

        # Grand Total Row (Daily Table)
        tot_row_num = start_row + len(pod_list)
        ws.row_dimensions[tot_row_num].height = 21
        tot_row_data = ['Grand Total']

        for d_idx in range(num_days):
            tot_col_let = get_column_letter(2 + d_idx * 5)
            aadh_col_let = get_column_letter(3 + d_idx * 5)
            man_col_let = get_column_letter(4 + d_idx * 5)
            red_col_let = get_column_letter(6 + d_idx * 5)

            tot_sum_f = f"=SUM({tot_col_let}3:{tot_col_let}{tot_row_num - 1})"
            aadh_sum_f = f"=SUM({aadh_col_let}3:{aadh_col_let}{tot_row_num - 1})"
            man_sum_f = f"=SUM({man_col_let}3:{man_col_let}{tot_row_num - 1})"
            comp_sum_f = f"=IFERROR({aadh_col_let}{tot_row_num}/{tot_col_let}{tot_row_num}, 1)"
            red_sum_f = f"=SUM({red_col_let}3:{red_col_let}{tot_row_num - 1})"

            tot_row_data.extend([tot_sum_f, aadh_sum_f, man_sum_f, comp_sum_f, red_sum_f])

        ws.append(tot_row_data)
        for c_idx in range(1, len(tot_row_data) + 1):
            cell = ws.cell(row=tot_row_num, column=c_idx)
            cell.font = FONT_APTOS_11_BOLD
            cell.fill = FILL_GRAND_TOTAL_DARK_25
            cell.border = BORDER_THIN_GRAY
            cell.alignment = ALIGN_LEFT if c_idx == 1 else ALIGN_CENTER
            if c_idx > 1 and (c_idx - 2) % 5 == 3:
                cell.number_format = FMT_PERCENT_INT
            elif c_idx > 1:
                cell.number_format = FMT_INT_COMMAS

        # --- MTD Table ---
        mtd_start = tot_row_num + 3
        ws.row_dimensions[mtd_start].height = 24
        ws.row_dimensions[mtd_start + 1].height = 20

        ws.append([])  # Spacing
        ws.append([])
        ws.append([label_header, 'MTD (D-1)', 'MTD (D-1)', 'MTD (D-1)', 'MTD (D-1)', 'MTD (D-1)'])
        ws.append([label_header, 'Total Onboarding', 'Aadhar', 'Manual', 'Compliance %', 'Red'])

        # Style MTD headers BEFORE merging
        for c in range(1, 7):
            cell_1 = ws.cell(row=mtd_start, column=c)
            cell_2 = ws.cell(row=mtd_start + 1, column=c)
            cell_1.border = BORDER_THIN_GRAY
            cell_2.border = BORDER_THIN_GRAY
            cell_1.alignment = ALIGN_CENTER
            cell_2.alignment = ALIGN_CENTER

            cell_1.font = FONT_APTOS_11_BOLD
            cell_1.fill = FILL_HEADER_BLUE_LIGHT_80
            cell_2.font = FONT_APTOS_11_BOLD
            if c == 5:
                cell_2.fill = FILL_YELLOW
            elif c == 6:
                cell_2.font = FONT_APTOS_11_BOLD_WHITE
                cell_2.fill = FILL_RED
            else:
                cell_2.fill = FILL_HEADER_BLUE_LIGHT_80

        ws.merge_cells(start_row=mtd_start, start_column=1, end_row=mtd_start + 1, end_column=1)
        ws.merge_cells(start_row=mtd_start, start_column=2, end_row=mtd_start, end_column=6)

        # MTD Data rows
        mtd_data_start = mtd_start + 2
        for p_idx, pod in enumerate(pod_list):
            cur_mtd_r = mtd_data_start + p_idx
            daily_r = 3 + p_idx
            ws.row_dimensions[cur_mtd_r].height = 19

            tot_parts = [f"{get_column_letter(2 + d * 5)}{daily_r}" for d in range(num_days)]
            aadh_parts = [f"{get_column_letter(3 + d * 5)}{daily_r}" for d in range(num_days)]
            man_parts = [f"{get_column_letter(4 + d * 5)}{daily_r}" for d in range(num_days)]
            red_parts = [f"{get_column_letter(6 + d * 5)}{daily_r}" for d in range(num_days)]

            ws.append([
                pod,
                f"=SUM({','.join(tot_parts)})",
                f"=SUM({','.join(aadh_parts)})",
                f"=SUM({','.join(man_parts)})",
                f"=IFERROR(C{cur_mtd_r}/B{cur_mtd_r}, 1)",
                f"=SUM({','.join(red_parts)})"
            ])

            for c_idx in range(1, 7):
                cell = ws.cell(row=cur_mtd_r, column=c_idx)
                cell.font = FONT_APTOS_11
                cell.border = BORDER_THIN_GRAY
                cell.alignment = ALIGN_LEFT if c_idx == 1 else ALIGN_CENTER
                if c_idx == 5:
                    cell.number_format = FMT_PERCENT_INT
                elif c_idx > 1:
                    cell.number_format = FMT_INT_COMMAS

        # MTD Grand Total
        mtd_tot_r = mtd_data_start + len(pod_list)
        ws.row_dimensions[mtd_tot_r].height = 21
        ws.append([
            'Grand Total',
            f"=SUM(B{mtd_data_start}:B{mtd_tot_r - 1})",
            f"=SUM(C{mtd_data_start}:C{mtd_tot_r - 1})",
            f"=SUM(D{mtd_data_start}:D{mtd_tot_r - 1})",
            f"=IFERROR(C{mtd_tot_r}/B{mtd_tot_r}, 1)",
            f"=SUM(F{mtd_data_start}:F{mtd_tot_r - 1})"
        ])
        for c_idx in range(1, 7):
            cell = ws.cell(row=mtd_tot_r, column=c_idx)
            cell.font = FONT_APTOS_11_BOLD
            cell.fill = FILL_GRAND_TOTAL_DARK_25
            cell.border = BORDER_THIN_GRAY
            cell.alignment = ALIGN_LEFT if c_idx == 1 else ALIGN_CENTER
            if c_idx == 5:
                cell.number_format = FMT_PERCENT_INT
            elif c_idx > 1:
                cell.number_format = FMT_INT_COMMAS

        # Column widths
        ws.column_dimensions['A'].width = 24
        for c_idx in range(2, len(tot_row_data) + 1):
            c_let = get_column_letter(c_idx)
            mod = (c_idx - 2) % 5
            ws.column_dimensions[c_let].width = 16 if mod == 4 else 15

    notify(0.35, "Writing Aadhaar KYC Compliance tabs (LM, HL, FM, HO-RO)...")
    populate_aadhar_summary(ws_lm, 'POD', LM_PODS)
    populate_aadhar_summary(ws_hl, 'Region', HL_REGIONS)
    populate_aadhar_summary(ws_fm, 'Region', FM_REGIONS)
    populate_aadhar_summary(ws_horo, 'Site', HORO_SITES)

    # -------------------------------------------------------------------------
    # SHEETS 5, 6, 7, 8: Manual HR Approval Breakdown
    # -------------------------------------------------------------------------
    def populate_manual_hr(ws, pod_list):
        ws.row_dimensions[1].height = 24
        ws.append(['Approved by', 'Manual cases', 'Red Case', 'total'])
        for cell in ws[1]:
            cell.font = FONT_APTOS_11_BOLD
            cell.fill = FILL_HEADER_BLUE_LIGHT_80
            cell.border = BORDER_THIN_GRAY
            cell.alignment = ALIGN_CENTER

        # Identify unique HR recruiters approving manual or red cases
        col_idx = {h.strip().lower(): i for i, h in enumerate(base_headers)}
        pod_col = col_idx.get('pod name', 10)
        hub_col = col_idx.get('hub name', 9)
        hr_col = col_idx.get('approved by', 18)
        aadh_col = col_idx.get('onboarded by aadhaar (y/n)', 19)
        auto_ver_col = col_idx.get('auto verification status', 23)
        hyp_status_col = col_idx.get('hyperverge status', 24)
        perm_addr_col = col_idx.get('permanent address', 30)

        pod_set = set(pod_list)
        hr_set = set()

        for row in base_rows:
            p_name = str(row[pod_col] or '').strip()
            h_name = str(row[hub_col] or '').strip()
            hr_name = str(row[hr_col] or '').strip()
            is_aadh_y = str(row[aadh_col] or '').strip().upper() == 'Y'
            auto_ver = str(row[auto_ver_col] or '').strip().lower() if auto_ver_col < len(row) else ''
            hyp_status = str(row[hyp_status_col] or '').strip().lower() if hyp_status_col < len(row) else ''
            perm_addr = str(row[perm_addr_col] or '').strip().lower() if perm_addr_col < len(row) else ''

            if 'hub migration' in perm_addr:
                continue

            if (p_name in pod_set or h_name in pod_set) and hr_name:
                is_manual = not is_aadh_y
                is_red = (auto_ver == 'red' or hyp_status == 'red' or auto_ver == 'failed')
                if is_manual or is_red:
                    hr_set.add(hr_name)

        sorted_hrs = sorted(list(hr_set))
        for h_idx, hr_name in enumerate(sorted_hrs, start=2):
            ws.row_dimensions[h_idx].height = 19
            ws.append([
                hr_name,
                f"=COUNTIFS('Onboarding Base'!$S:$S, $A{h_idx}, 'Onboarding Base'!$T:$T, \"N\", 'Onboarding Base'!$D:$D, \"Onboarded*\", 'Onboarding Base'!$AE:$AE, \"<>*hub migration*\")",
                f"=COUNTIFS('Onboarding Base'!$S:$S, $A{h_idx}, 'Onboarding Base'!$X:$X, \"red\", 'Onboarding Base'!$D:$D, \"Onboarded*\", 'Onboarding Base'!$AE:$AE, \"<>*hub migration*\")",
                f"=B{h_idx}+C{h_idx}"
            ])
            for c in range(1, 5):
                cell = ws.cell(row=h_idx, column=c)
                cell.font = FONT_APTOS_11
                cell.border = BORDER_THIN_GRAY
                cell.alignment = ALIGN_LEFT if c == 1 else ALIGN_CENTER
                if c > 1:
                    cell.number_format = FMT_INT_COMMAS

        # Grand Total
        tot_r = len(sorted_hrs) + 2
        ws.row_dimensions[tot_r].height = 21
        end_r = tot_r - 1
        ws.append([
            'Grand Total',
            f"=SUM(B2:B{end_r})" if len(sorted_hrs) > 0 else 0,
            f"=SUM(C2:C{end_r})" if len(sorted_hrs) > 0 else 0,
            f"=SUM(D2:D{end_r})" if len(sorted_hrs) > 0 else 0
        ])
        for c in range(1, 5):
            cell = ws.cell(row=tot_r, column=c)
            cell.font = FONT_APTOS_11_BOLD
            cell.fill = FILL_GRAND_TOTAL_DARK_25
            cell.border = BORDER_THIN_GRAY
            cell.alignment = ALIGN_LEFT if c == 1 else ALIGN_CENTER
            if c > 1:
                cell.number_format = FMT_INT_COMMAS

        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 18

    notify(0.70, "Writing Manual Approvals & Recruiter attribution tabs...")
    populate_manual_hr(ws_man_lm, LM_PODS)
    populate_manual_hr(ws_man_hl, HL_REGIONS)
    populate_manual_hr(ws_man_fm, FM_REGIONS)
    populate_manual_hr(ws_man_horo, HORO_SITES)

    notify(0.96, "Saving Onboarding Report workbook to disk...")
    wb.save(output_path)
    wb.close()
    notify(1.00, "Onboarding Report workbook generated.")


# ==============================================================================
# PUNCH ADHERENCE REPORT GENERATOR
# ==============================================================================

def build_punch_adherence_report(
    output_path: str,
    date_list: List[str],
    workforce_records: List[Dict[str, Any]],
    base_hub_combos: List[Tuple[str, str, str]],
    progress_callback: Optional[Callable[[float, str], None]] = None
):
    """
    Constructs the 7-sheet Punch Adherence Report with exact Aptos Narrow formatting,
    live COUNTIFS & VLOOKUP formulas, Red/Amber/Green headers, and Base facility summaries.
    """
    def notify(pct: float, msg: str):
        if progress_callback:
            progress_callback(pct, msg)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # 1. Instantiate worksheets in frontmost order
    notify(0.05, "Creating worksheets in frontmost order...")
    ws_lm = wb.create_sheet(title='Attendance Adherence - LM')
    ws_hl = wb.create_sheet(title='Attendance Adherence - HL')
    ws_fm = wb.create_sheet(title='Attendance Adherence - FM')
    ws_horo = wb.create_sheet(title='Attendance Adherence - HO-RO')
    ws_att = wb.create_sheet(title='Attendance Base')
    ws_base = wb.create_sheet(title='Base')
    ws_punched = wb.create_sheet(title='Punched Base')

    num_days = len(date_list)
    total_wf = len(workforce_records)

    # -------------------------------------------------------------------------
    # Helper: Populate Summary Sheets (LM, HL, FM, HO-RO)
    # -------------------------------------------------------------------------
    def populate_summary_sheet(ws, label_header: str, entity_list: List[str]):
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 20

        r1 = [label_header, 'Onroll HC', 'Offroll HC', 'Total HC', 'Red', 'Amber', 'Green', 'Compliance %']
        r2 = [None, None, None, None, 'No Punches(30+D)', 'Punches Received (4-29D)', 'Recent Punches(3D)', None]
        ws.append(r1)
        ws.append(r2)

        ws.merge_cells('A1:A2')
        ws.merge_cells('B1:B2')
        ws.merge_cells('C1:C2')
        ws.merge_cells('D1:D2')
        ws.merge_cells('H1:H2')

        # Style Header Row 1
        for c in range(1, 9):
            cell = ws.cell(row=1, column=c)
            cell.font = FONT_APTOS_9_BOLD
            cell.alignment = ALIGN_CENTER_WRAP
            cell.border = BORDER_THIN_GRAY
            if c == 5:
                cell.fill = FILL_PUNCH_RED
            elif c == 6:
                cell.fill = FILL_PUNCH_AMBER
            elif c == 7:
                cell.fill = FILL_PUNCH_GREEN
            else:
                cell.fill = FILL_PUNCH_HDR

        # Style Header Row 2
        for c in range(1, 9):
            cell = ws.cell(row=2, column=c)
            cell.font = FONT_APTOS_9_BOLD
            cell.alignment = ALIGN_CENTER_WRAP
            cell.border = BORDER_THIN_GRAY
            if c in (5, 6, 7):
                cell.fill = FILL_PUNCH_HDR

        # Data Rows
        start_r = 3
        for idx, entity in enumerate(entity_list):
            cur_r = start_r + idx
            ws.row_dimensions[cur_r].height = 19

            f_on = f"=COUNTIFS('Punched Base'!$D:$D, $A{cur_r}, 'Punched Base'!$C:$C, \"On-Roll\")"
            f_off = f"=COUNTIFS('Punched Base'!$D:$D, $A{cur_r}, 'Punched Base'!$C:$C, \"Off-Roll\")"
            f_tot = f"=B{cur_r}+C{cur_r}"
            f_red = f"=COUNTIFS('Punched Base'!$D:$D, $A{cur_r}, 'Punched Base'!$H:$H, \"Red\")"
            f_amb = f"=COUNTIFS('Punched Base'!$D:$D, $A{cur_r}, 'Punched Base'!$H:$H, \"Amber\")"
            f_grn = f"=COUNTIFS('Punched Base'!$D:$D, $A{cur_r}, 'Punched Base'!$H:$H, \"Green\")"
            f_comp = f"=IFERROR(G{cur_r}/D{cur_r},1)"

            row_data = [entity, f_on, f_off, f_tot, f_red, f_amb, f_grn, f_comp]
            ws.append(row_data)

            for c in range(1, 9):
                cell = ws.cell(row=cur_r, column=c)
                cell.font = FONT_APTOS_11
                cell.border = BORDER_THIN_GRAY
                if c == 1:
                    cell.alignment = ALIGN_LEFT
                else:
                    cell.alignment = ALIGN_CENTER
                    if c == 8:
                        cell.number_format = FMT_PERCENT_INT
                    else:
                        cell.number_format = FMT_INT_COMMAS

        # Grand Total Row
        tot_r = start_r + len(entity_list)
        ws.row_dimensions[tot_r].height = 21
        end_r = tot_r - 1

        tot_vals = [
            'Grand Total',
            f"=SUM(B3:B{end_r})",
            f"=SUM(C3:C{end_r})",
            f"=SUM(D3:D{end_r})",
            f"=SUM(E3:E{end_r})",
            f"=SUM(F3:F{end_r})",
            f"=SUM(G3:G{end_r})",
            f"=IFERROR(G{tot_r}/D{tot_r},1)"
        ]
        ws.append(tot_vals)

        for c in range(1, 9):
            cell = ws.cell(row=tot_r, column=c)
            cell.font = FONT_APTOS_11_BOLD
            cell.fill = FILL_PUNCH_TOTAL
            cell.border = BORDER_THIN_GRAY
            if c == 1:
                cell.alignment = ALIGN_LEFT
            else:
                cell.alignment = ALIGN_CENTER
                if c == 8:
                    cell.number_format = FMT_PERCENT_INT
                else:
                    cell.number_format = FMT_INT_COMMAS

        ws.column_dimensions['A'].width = 22.0
        for c_let in ('B', 'C', 'D', 'E', 'F', 'G', 'H'):
            ws.column_dimensions[c_let].width = 13.0

    # -------------------------------------------------------------------------
    # SHEET 5: Attendance Base
    # -------------------------------------------------------------------------
    notify(0.15, f"Writing Attendance Base ({total_wf:,} rows)...")
    ws_att.row_dimensions[1].height = 24
    formatted_date_headers = [format_date_header(d) for d in date_list]
    att_h1 = ['ID', 'Name', 'Type', 'POD', 'Registered Facility', 'Department', 'Role'] + formatted_date_headers + ['Punches in last 30days', 'Punches in last 3days', 'Color Status']
    ws_att.append(att_h1)

    for c in range(1, len(att_h1) + 1):
        cell = ws_att.cell(row=1, column=c)
        cell.font = FONT_APTOS_9_BOLD
        cell.fill = FILL_PUNCH_HDR
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY

    # Blank Row 2 matching reference
    ws_att.append([None] * len(att_h1))
    for c in range(1, len(att_h1) + 1):
        cell = ws_att.cell(row=2, column=c)
        cell.border = BORDER_THIN_GRAY

    first_date_col = get_column_letter(8)
    last_date_col = get_column_letter(7 + num_days)
    third_last_col = get_column_letter(max(8, 7 + num_days - 2))
    col_30d_let = get_column_letter(8 + num_days)
    col_3d_let = get_column_letter(9 + num_days)

    for idx, emp in enumerate(workforce_records):
        cur_r = 3 + idx
        f_30d = f'=IF(COUNTIF({first_date_col}{cur_r}:{last_date_col}{cur_r},"P")>0,"Y","N")'
        f_3d = f'=IF(COUNTIF({third_last_col}{cur_r}:{last_date_col}{cur_r},"P")>0,"Y","N")'
        f_col = f'=IF({col_3d_let}{cur_r}="Y","Green",IF({col_30d_let}{cur_r}="Y","Amber","Red"))'

        d_punches = [emp['punches'].get(d, 'A') for d in date_list]
        row_vals = [
            emp['id'], emp['name'], emp['type'], emp['pod'], emp['facility'],
            emp['department'], emp['role']
        ] + d_punches + [f_30d, f_3d, f_col]
        ws_att.append(row_vals)

        ws_att.cell(row=cur_r, column=1).number_format = FMT_TEXT

        if idx > 0 and idx % 4000 == 0:
            notify(0.15 + (idx / total_wf) * 0.35, f"Writing Attendance Base: {idx:,} / {total_wf:,} rows...")

    for row in ws_att.iter_rows(min_row=3, max_row=2 + total_wf):
        for idx, cell in enumerate(row):
            cell.font = FONT_APTOS_11
            cell.border = BORDER_THIN_GRAY
            if idx in (0, 1, 3, 4, 5, 6):
                cell.alignment = ALIGN_LEFT
            else:
                cell.alignment = ALIGN_CENTER

    ws_att.column_dimensions['A'].width = 16.0
    ws_att.column_dimensions['B'].width = 24.0
    ws_att.column_dimensions['C'].width = 14.0
    ws_att.column_dimensions['D'].width = 22.0
    ws_att.column_dimensions['E'].width = 24.0
    ws_att.column_dimensions['F'].width = 18.0
    ws_att.column_dimensions['G'].width = 18.0
    for c_idx in range(8, len(att_h1) + 1):
        ws_att.column_dimensions[get_column_letter(c_idx)].width = 13.0

    # -------------------------------------------------------------------------
    # SHEET 7: Punched Base
    # -------------------------------------------------------------------------
    notify(0.55, f"Writing Punched Base ({total_wf:,} rows)...")
    ws_punched.row_dimensions[1].height = 24
    pb_h1 = ['ID', 'Name', 'Type', 'POD', 'Hub Name', 'Punches in last 30days', 'Punches in last 3days', 'Color Status']
    ws_punched.append(pb_h1)

    for c in range(1, len(pb_h1) + 1):
        cell = ws_punched.cell(row=1, column=c)
        cell.font = FONT_APTOS_9_BOLD
        cell.fill = FILL_PUNCH_HDR
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN_GRAY

    col_30d_idx = 8 + num_days
    col_3d_idx = 9 + num_days

    for idx, emp in enumerate(workforce_records):
        cur_r = 2 + idx
        f_v_30d = f"=VLOOKUP(A{cur_r},'Attendance Base'!$A:$T,{col_30d_idx},FALSE)"
        f_v_3d = f"=VLOOKUP(A{cur_r},'Attendance Base'!$A:$T,{col_3d_idx},FALSE)"
        f_v_col = f'=IF(G{cur_r}="Y","Green",IF(F{cur_r}="Y","Amber","Red"))'

        row_vals = [emp['id'], emp['name'], emp['type'], emp['pod'], emp['facility'], f_v_30d, f_v_3d, f_v_col]
        ws_punched.append(row_vals)
        ws_punched.cell(row=cur_r, column=1).number_format = FMT_TEXT

    for row in ws_punched.iter_rows(min_row=2, max_row=1 + total_wf):
        for idx, cell in enumerate(row):
            cell.font = FONT_APTOS_11
            cell.border = BORDER_THIN_GRAY
            if idx in (0, 1, 3, 4):
                cell.alignment = ALIGN_LEFT
            else:
                cell.alignment = ALIGN_CENTER

    ws_punched.column_dimensions['A'].width = 16.0
    ws_punched.column_dimensions['B'].width = 24.0
    ws_punched.column_dimensions['C'].width = 14.0
    ws_punched.column_dimensions['D'].width = 22.0
    ws_punched.column_dimensions['E'].width = 24.0
    ws_punched.column_dimensions['F'].width = 14.0
    ws_punched.column_dimensions['G'].width = 14.0
    ws_punched.column_dimensions['H'].width = 14.0

    # -------------------------------------------------------------------------
    # SHEET 6: Base (Facility Summary)
    # -------------------------------------------------------------------------
    notify(0.70, f"Writing Base ({len(base_hub_combos):,} facilities)...")
    ws_base.row_dimensions[1].height = 24
    base_h1 = ['Hub Name', 'POD Name', 'Veticle', 'Onroll HC', 'Offroll HC', 'Total HC', 'Red', 'Amber', 'Green', 'Compliance %']
    ws_base.append(base_h1)

    for c in range(1, len(base_h1) + 1):
        cell = ws_base.cell(row=1, column=c)
        cell.font = FONT_APTOS_9_BOLD
        cell.alignment = ALIGN_CENTER_WRAP
        cell.border = BORDER_THIN_GRAY
        if c == 7:
            cell.fill = FILL_PUNCH_RED
        elif c == 8:
            cell.fill = FILL_PUNCH_AMBER
        elif c == 9:
            cell.fill = FILL_PUNCH_GREEN
        else:
            cell.fill = FILL_PUNCH_HDR

    for idx, (hub, pod, vehicle) in enumerate(base_hub_combos):
        cur_r = 2 + idx
        f_on = f"=COUNTIFS('Punched Base'!$E:$E, $A{cur_r}, 'Punched Base'!$D:$D, $B{cur_r}, 'Punched Base'!$C:$C, \"On-Roll\")"
        f_off = f"=COUNTIFS('Punched Base'!$E:$E, $A{cur_r}, 'Punched Base'!$D:$D, $B{cur_r}, 'Punched Base'!$C:$C, \"Off-Roll\")"
        f_tot = f"=D{cur_r}+E{cur_r}"
        f_red = f"=COUNTIFS('Punched Base'!$E:$E, $A{cur_r}, 'Punched Base'!$D:$D, $B{cur_r}, 'Punched Base'!$H:$H, \"Red\")"
        f_amb = f"=COUNTIFS('Punched Base'!$E:$E, $A{cur_r}, 'Punched Base'!$D:$D, $B{cur_r}, 'Punched Base'!$H:$H, \"Amber\")"
        f_grn = f"=COUNTIFS('Punched Base'!$E:$E, $A{cur_r}, 'Punched Base'!$D:$D, $B{cur_r}, 'Punched Base'!$H:$H, \"Green\")"
        f_comp = f"=IFERROR(I{cur_r}/F{cur_r},1)"

        row_vals = [hub, pod, vehicle, f_on, f_off, f_tot, f_red, f_amb, f_grn, f_comp]
        ws_base.append(row_vals)

    for row in ws_base.iter_rows(min_row=2, max_row=1 + len(base_hub_combos)):
        for idx, cell in enumerate(row):
            cell.font = FONT_APTOS_11
            cell.border = BORDER_THIN_GRAY
            if idx in (0, 1, 2):
                cell.alignment = ALIGN_LEFT
            else:
                cell.alignment = ALIGN_CENTER
                if idx == 9:
                    cell.number_format = FMT_PERCENT_INT
                else:
                    cell.number_format = FMT_INT_COMMAS

    ws_base.column_dimensions['A'].width = 22.0
    for c_let in ('B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J'):
        ws_base.column_dimensions[c_let].width = 13.0

    # -------------------------------------------------------------------------
    # Populate Summary Tabs (Frontmost 1 to 4)
    # -------------------------------------------------------------------------
    notify(0.85, "Writing Executive Summary tabs (LM, HL, FM, HO-RO)...")
    populate_summary_sheet(ws_lm, 'POD', PUNCH_LM_PODS)
    populate_summary_sheet(ws_hl, 'Region', PUNCH_HL_REGIONS)
    populate_summary_sheet(ws_fm, 'Region', PUNCH_FM_REGIONS)
    populate_summary_sheet(ws_horo, 'Site', PUNCH_HORO_SITES)

    notify(0.96, "Saving Punch Adherence Report workbook to disk...")
    wb.save(output_path)
    wb.close()
    notify(1.00, "Punch Adherence Report generated successfully!")

