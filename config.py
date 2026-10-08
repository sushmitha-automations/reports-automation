"""
Configuration and Master Reference Constants for Automation Reports.
Contains exact POD lists, region mappings, styling palettes, fonts, borders, and number formats.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==============================================================================
# OPERATIONAL VERTICALS & ENTITY LISTS
# ==============================================================================

LM_PODS = [
    "POD_MP/RJ", "POD_ROTN", "POD_BLR", "POD_JK/HP/HR/PB/UK", "POD_HYD",
    "POD_BOM/PNQ", "POD_ROMH", "POD_CHN", "POD_UP", "POD_AP/TS", "POD_NE",
    "POD_BH/JH", "POD_ROK/KL", "POD_GJ", "POD_NCR/UP_West", "POD_CCU/ROWB", "POD_CG/OD"
]

HL_REGIONS = ["HL_North", "HL_West", "HL_East", "HL_South"]

FM_REGIONS = ["FM_East", "FM_South", "FM_West", "FM_North"]

HORO_SITES = [
    "JAIPUR - RO", "PATNA - RO", "Bangalore - RO", "MUMBAI - RO", "GURUGRAM - RO",
    "KOLKATA - RO", "PUNE - RO", "HYDERABAD - RO", "CHENNAI - RO", "LUCKNOW - RO",
    "Nagpur RO", "Jaipur RO", "Bangalore - HO", "CHANDIGARH - RO", "BHUBANESWAR - RO"
]

# Validation Reference Rules (Chroma Status vs Valinor Attendance Status)
REF_VAL_RULES = [
    ("P", "P", "MATCH"),
    ("P", "A", "P in Chroma, A in Valinor"),
    ("A", "P", "A/LWP in Chroma, P in Valinor"),
    ("A", "A", "MATCH"),
    ("LWP", "P", "A/LWP in Chroma, P in Valinor"),
    ("LWP", "A", "MATCH"),
    ("NA", "P", "NA in Chroma, P in Valinor"),
    ("NA", "A", "MATCH"),
    ("WO", "P", "WO/PL/HD in Chroma, P in Valinor"),
    ("WO", "A", "WO/PL/HD in Chroma, A in Valinor"),
    ("PL", "P", "WO/PL/HD in Chroma, P in Valinor"),
    ("PL", "A", "WO/PL/HD in Chroma, A in Valinor"),
    ("HD", "P", "MATCH"),
    ("HD", "A", "WO/PL/HD in Chroma, A in Valinor"),
    ("NPP", "P", "MATCH"),
    ("NPP", "A", "P in Chroma, A in Valinor")
]

# Role Reference Rules
REF_ROLE_RULES = [
    ("DE", "DE - BIKER"),
    ("DE", "DE - FOUR WHEELER"),
    ("DE", "DE - THREE WHEELER"),
    ("DE", "DE - VAN DRIVER"),
    ("DE", "DE - ELECTRIC VEHICLE"),
    ("DE", "TEAM LEADER"),
    ("TL", "TEAM LEADER"),
    ("TL", "TL"),
    ("TL", "FIELD TL"),
    ("LM_DE", "DE"),
    ("FM_DE", "DE"),
    ("HL_DE", "DE")
]

# ==============================================================================
# STYLING PALETTES: EXACT REPORT MATCHING (NO OVER-STYLING)
# ==============================================================================

# Borders
BORDER_THIN_GRAY = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

BORDER_THIN_BLACK = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)

BORDER_DOUBLE_BOTTOM = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='double', color='000000')
)

# Colors & Fills
FILL_DARK_BLUE = PatternFill(start_color='1F497D', end_color='1F497D', fill_type='solid')      # Primary dark blue headers
FILL_LIGHT_BLUE = PatternFill(start_color='DBE9F7', end_color='DBE9F7', fill_type='solid')     # Onboarding headers
FILL_SOFT_BLUE_HDR = PatternFill(start_color='B8CCE4', end_color='B8CCE4', fill_type='solid')  # Attendance check sub-headers
FILL_YELLOW = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')         # Compliance % header
FILL_RED = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')            # Red cases header
FILL_LIGHT_GRAY = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')     # Grand total rows

# Fonts
FONT_CALIBRI_11 = Font(name='Calibri', size=11, bold=False, color='000000')
FONT_CALIBRI_11_BOLD = Font(name='Calibri', size=11, bold=True, color='000000')
FONT_CALIBRI_11_BOLD_WHITE = Font(name='Calibri', size=11, bold=True, color='FFFFFF')

FONT_APTOS_11 = Font(name='Aptos Narrow', size=11, bold=False, color='000000')
FONT_APTOS_11_BOLD = Font(name='Aptos Narrow', size=11, bold=True, color='000000')
FONT_APTOS_11_BOLD_WHITE = Font(name='Aptos Narrow', size=11, bold=True, color='FFFFFF')

# Alignments
ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_CENTER_WRAP = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_LEFT = Alignment(horizontal='left', vertical='center')
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center')

# Number Formats
FMT_TEXT = '@'
FMT_PERCENT_2DEC = '0.00%'
FMT_PERCENT_INT = '0%'
FMT_INT_COMMAS = '#,##0'
FMT_DATE_ISO = 'YYYY-MM-DD'
FMT_DATE_HEADER = 'DD-MMM-YY'
